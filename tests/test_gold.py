"""Regression checks for the deterministic service boundary."""
from __future__ import annotations

import os
import base64
import json
from io import BytesIO
import tempfile
import zipfile
from pathlib import Path

import pandas as pd
import pytest
from fastapi.testclient import TestClient

from analysis_service import analyze_dataframe, generate_docx
from api import app
from data_parser import (
    DataParserError,
    parse_csv_buffer,
    parse_docx_buffer,
    parse_excel_buffer,
    parse_image,
    parse_pdf,
    parse_tabular_text,
    parse_uploaded_file,
)
from financial_engine import analyze_financial_dataframe
from handle import build_breakdown


def test_research_rejects_nonfinite_designs_and_reports_exclusions() -> None:
    from stats_engine import correlation, linear_regression, paired_ttest

    with pytest.raises(ValueError, match="non-constant"):
        correlation([1, 1, 1], [1, 2, 3])
    with pytest.raises(ValueError, match="non-constant"):
        linear_regression([1, 1, 1], [1, 2, 3])
    with pytest.raises(ValueError, match="variable paired differences"):
        paired_ttest([1, 2, 3], [2, 3, 4])

    engine = analyze_dataframe(
        pd.DataFrame({"Treatment": ["A", "A", "B", "B"], "Value": [1, "bad", 4, 5]}),
        "Treatment",
        "Value",
    )
    quality = engine["ingested"]["data_quality"]
    assert quality["input_rows"] == 4
    assert quality["analyzed_rows"] == 3
    assert quality["exclusion_reasons"] == {"non_numeric_metric": 1}


def test_low_quality_data_still_returns_a_result_without_refusal() -> None:
    from handle import handle_analyze

    bad_ingested = {
        "format": "long",
        "factor": "Treatment",
        "outcome": "Moisture",
        "groups": [
            {"name": "A", "values": [101, 102]},
            {"name": "B", "values": [20, 21]},
        ],
    }
    import handle as handle_module

    original_ingest = handle_module.ingest_text
    handle_module.ingest_text = lambda _: bad_ingested
    try:
        result = handle_analyze({"text": "placeholder"})
    finally:
        handle_module.ingest_text = original_ingest
    assert result["ok"] is True
    assert result.get("refused") is not True
    assert result["result"]["test"] in {"student-t", "welch-t", "one-way anova"}


def test_financial_analysis_rejects_invalid_numbers_and_discloses_proxies() -> None:
    with pytest.raises(ValueError, match="non-numeric values"):
        analyze_financial_dataframe(pd.DataFrame({
            "Month": ["2026-01", "2026-02"],
            "Revenue": [100, "bad"],
            "Operating Expenses": [40, 45],
        }))

    result = analyze_financial_dataframe(pd.DataFrame({
        "Month": ["2026-01", "2026-02"],
        "Revenue": [100, 110],
        "Operating Expenses": [40, 45],
        "Payroll": [10, 10],
    }))
    assert result["cash_flow_status"] == "ebitda_proxy"
    assert any(warning["type"] == "overlapping_expense_columns" for warning in result["diagnostics"]["warnings"])

    reported = analyze_financial_dataframe(pd.DataFrame({
        "Month": ["2026-01", "2026-02"],
        "Revenue": [100, 110],
        "Operating Expenses": [40, 45],
        "Operating Cash Flow": [55, 60],
    }))
    assert reported["cash_flow_status"] == "reported"
    assert reported["periods"][-1]["net_operating_cash_flow"] == 60.0


def test_parser_delimiters() -> None:
    assert list(parse_csv_buffer(b"Group,Score\nA,1\nB,2\n").columns) == ["Group", "Score"]
    assert parse_tabular_text("Group\tScore\nA\t1").shape == (1, 2)
    assert parse_tabular_text("Group Score\nA 1").shape == (1, 2)


def test_financial_engine_returns_kpis_diagnostics_narrative_and_chart() -> None:
    frame = pd.DataFrame({
        "Month": ["2026-01", "2026-02", "2026-04"],
        "Revenue": [100000, 110000, 121000],
        "Operating Expenses": [40000, 42000, 140000],
        "Cash Reserves": [200000, 160000, 40000],
    })
    result = analyze_financial_dataframe(frame)

    assert result["kpis"]["net_revenue"] == 121000.0
    assert result["kpis"]["ebitda"] == -19000.0
    assert abs(result["kpis"]["runway_months"] - 0.5405405405) < 1e-9
    assert result["kpis"]["revenue_trend"]["r_squared"] is not None
    assert any(warning["type"] == "mismatched_date_sequence" for warning in result["diagnostics"]["warnings"])
    assert any(
        warning["type"] == "negative_cash_flow" and warning["severity"] == "critical"
        for warning in result["diagnostics"]["warnings"]
    )
    assert "Financial Performance Summary:" in result["executive_summary"]
    assert result["chart_base64"].startswith("iVBORw0KGgo")


def test_financial_engine_accepts_currency_suffixed_headers() -> None:
    frame = pd.DataFrame({
        "Month": ["Jan", "Feb", "Mar"],
        "Department": ["Operations"] * 3,
        "Revenue_USD": [45000, 48000, 52000],
        "Operating_Expenses_USD": [12000, 11500, 13000],
        "Payroll_USD": [18000, 18000, 18500],
        "Marketing_Spend_USD": [5000, 4500, 6000],
        "Active_Clients": [120, 125, 135],
    })
    result = analyze_financial_dataframe(frame)

    assert result["kpis"]["net_revenue"] == 52000.0
    assert abs(result["kpis"]["mean_monthly_expenses"] - 12166.666666666666) < 1e-9


def test_financial_docx_includes_chart_images() -> None:
    frame = pd.DataFrame({
        "Month": ["2026-01", "2026-02", "2026-03"],
        "Revenue": [100000, 110000, 120000],
        "Operating Expenses": [35000, 40000, 45000],
        "Cash Reserves": [200000, 180000, 150000],
    })
    engine = analyze_financial_dataframe(frame)
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "financial-results.docx"
        generate_docx(engine, path)
        with zipfile.ZipFile(path) as archive:
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
            assert media
            chart_payloads = [archive.read(name) for name in media]
            assert any(payload.startswith(b"\x89PNG\r\n\x1a\n") for payload in chart_payloads)


def test_financial_summary_uses_totals_and_missing_cash_sets_na_runway() -> None:
    frame = pd.DataFrame({
        "Month": ["2026-01", "2026-02", "2026-03"],
        "Revenue": [100, 200, 300],
        "Operating Expenses": [40, 50, 60],
        "Payroll": [10, 10, 10],
    })
    result = analyze_financial_dataframe(frame)
    assert result["template_context"]["total_net_revenue"] == 600.0
    assert result["kpis"]["total_net_revenue"] == 600.0
    assert result["kpis"]["cumulative_ebitda"] == 600.0 - (40 + 50 + 60)
    assert result["kpis"]["runway_months"] == "N/A"
    assert not any(warning.get("type") == "cash_runway_warning" for warning in result["diagnostics"]["warnings"])


def test_financial_api_accepts_json_rows() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    response = TestClient(app).post(
        "/api/v1/financial-analysis",
        json={"rows": [
            {"Month": "2026-01", "Revenue": 100, "Operating Expenses": 40, "Cash Reserves": 200},
            {"Month": "2026-02", "Revenue": 110, "Operating Expenses": 45, "Cash Reserves": 155},
        ]},
        headers={"X-Rowfirst-Id": "demo-id", "X-Rowfirst-Secret-Key": "demo-secret"},
    )
    assert response.status_code == 200
    assert response.json()["analysis_type"] == "executive_financial"


def test_production_health_routes_are_available() -> None:
    client = TestClient(app)
    assert client.get("/health").status_code == 200
    assert client.get("/api/healthz").status_code == 200
    assert client.get("/readyz").status_code in {200, 503}


def test_financial_api_accepts_word_upload() -> None:
    from docx import Document

    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    document = Document()
    table = document.add_table(rows=1, cols=4)
    for cell, header in zip(table.rows[0].cells, ("Month", "Revenue", "Operating Expenses", "Cash Reserves")):
        cell.text = header
    for values in (("2026-01", "100", "40", "200"), ("2026-02", "110", "45", "155")):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    document_buffer = BytesIO()
    document.save(document_buffer)

    response = TestClient(app).post(
        "/api/v1/financial-analysis",
        files={"file": ("ledger.docx", document_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers={"X-Rowfirst-Id": "demo-id", "X-Rowfirst-Secret-Key": "demo-secret"},
    )
    assert response.status_code == 200
    assert response.json()["kpis"]["net_revenue"] == 110.0


def test_research_file_uploads_work_without_manual_column_selection() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    headers = {"X-Rowfirst-Id": "demo-id", "X-Rowfirst-Secret-Key": "demo-secret"}
    client = TestClient(app)
    csv_data = b"Treatment,Value\nA,1\nA,2\nB,4\nB,5\n"
    response = client.post(
        "/api/v1/analyze",
        files={"file": ("study.csv", csv_data, "text/csv")},
        data={"response_format": "json"},
        headers=headers,
    )
    assert response.status_code == 200
    assert response.json()["metric"] == "Value"

    excel_buffer = BytesIO()
    pd.DataFrame({"Treatment": ["A", "A", "B", "B"], "Value": [1, 2, 4, 5]}).to_excel(excel_buffer, index=False)
    response = client.post(
        "/api/v1/analyze",
        files={"file": ("study.xlsx", excel_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"response_format": "json"},
        headers=headers,
    )
    assert response.status_code == 200
    assert response.json()["metric"] == "Value"


def test_batch_analysis_processes_items_sequentially_with_isolated_errors() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    headers = {"X-Rowfirst-Id": "demo-id", "X-Rowfirst-Secret-Key": "demo-secret"}
    response = TestClient(app).post(
        "/api/v1/batch-analyze",
        json={"items": [
            {"name": "study-a", "raw_text": "Treatment,Value\nA,1\nA,2\nB,4\nB,5\n"},
            {"name": "bad-item", "raw_text": "not a table"},
            {"name": "study-b", "raw_text": "Treatment,Value\nA,2\nA,3\nB,5\nB,6\n"},
        ]},
        headers=headers,
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["batch"] is True
    assert payload["count"] == 3
    assert [item["status"] for item in payload["results"]] == ["success", "error", "success"]


def test_batch_analysis_accepts_multiple_uploaded_files() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    headers = {"X-Rowfirst-Id": "demo-id", "X-Rowfirst-Secret-Key": "demo-secret"}
    csv_a = b"Treatment,Value\nA,1\nA,2\nB,4\nB,5\n"
    csv_b = b"Treatment,Value\nA,2\nA,3\nB,5\nB,6\n"
    response = TestClient(app).post(
        "/api/v1/batch-analyze",
        files=[
            ("files", ("study-a.csv", csv_a, "text/csv")),
            ("files", ("study-b.csv", csv_b, "text/csv")),
        ],
        headers=headers,
    )
    assert response.status_code == 200
    assert [item["name"] for item in response.json()["results"]] == ["study-a.csv", "study-b.csv"]


def test_batch_download_returns_zip_reports_and_manifest() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    response = TestClient(app).post(
        "/api/v1/batch-analyze/download",
        json={"items": [
            {"name": "study-a", "raw_text": "Treatment,Value\nA,1\nA,2\nB,4\nB,5\n"},
            {"name": "study-b", "raw_text": "Treatment,Value\nA,2\nA,3\nB,5\nB,6\n"},
        ]},
        headers={"X-Rowfirst-Id": "demo-id", "X-Rowfirst-Secret-Key": "demo-secret"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/zip")
    with zipfile.ZipFile(BytesIO(response.content)) as archive:
        names = archive.namelist()
        assert "manifest.json" in names
        assert "001_study-a.docx" in names
        assert "002_study-b.docx" in names
        manifest = json.loads(archive.read("manifest.json"))
        assert [item["status"] for item in manifest["items"]] == ["success", "success"]

def test_telegram_polling_requires_token(monkeypatch) -> None:
    import api as api_module

    monkeypatch.delenv("TELEGRAM_BOT_TOKEN", raising=False)
    monkeypatch.delenv("BOT_TOKEN", raising=False)
    api_module._telegram_started = False
    api_module._start_telegram_polling()
    assert api_module._telegram_started is False


def test_labeled_text_and_analysis() -> None:
    frame = parse_tabular_text("A: 1, 2, 3\nB: 4, 5, 6")
    engine = analyze_dataframe(frame, "Factor", "Metric")
    assert engine["ok"] is True
    assert engine["factor"] == "Factor"
    assert engine["metric"] == "Metric"
    assert "n caveat" not in engine["breakdown"]
    assert "Sample-size caveat:" in engine["breakdown"]


def test_api_inference_skips_time_and_index_columns_and_requires_measurable_outcomes() -> None:
    from api import _infer_columns, _numeric_metric_columns

    frame = pd.DataFrame({
        "Record_ID": [f"R{index:03d}" for index in range(20)],
        "Teaching_Method": ["A"] * 10 + ["B"] * 10,
        "Week": list(range(1, 21)),
        "Month": [1] * 20,
        "Exam_Score": list(range(60, 80)),
    })
    assert _infer_columns(frame) == ("Teaching_Method", "Exam_Score")
    assert _numeric_metric_columns(frame, "Teaching_Method") == ["Exam_Score"]

    frame["Exam_Score"] = 1
    try:
        _infer_columns(frame)
    except DataParserError as exc:
        assert "No measurable outcome" in str(exc)
    else:
        raise AssertionError("constant or non-measurable outcomes should not be inferred")


def test_api_metric_picker_ignores_time_fields_for_finance_batch_analysis() -> None:
    from api import _numeric_metric_columns

    frame = pd.DataFrame({
        "Branch": ["A", "A", "B", "B", "A", "A", "B", "B"],
        "Week": [1, 2, 1, 2, 1, 2, 1, 2],
        "Month": ["Jan", "Jan", "Feb", "Feb", "Mar", "Mar", "Apr", "Apr"],
        "Deposits_NGN": [100, 120, 90, 110, 130, 150, 95, 115],
        "Withdrawal_Count": [10, 12, 9, 11, 13, 15, 10, 12],
        "NPS": [42, 50, 38, 41, 48, 53, 40, 44],
        "Error_Tickets": [5, 7, 6, 8, 9, 11, 8, 10],
    })
    assert _numeric_metric_columns(frame, "Branch") == [
        "Deposits_NGN",
        "Withdrawal_Count",
        "NPS",
        "Error_Tickets",
    ]


def test_api_auto_infers_named_group_columns_without_asking() -> None:
    from api import _design_gate, _infer_factor_column

    frame = pd.DataFrame({
        "Product": ["A", "A", "B", "B", "C", "C", "A", "A", "B", "B", "C", "C"],
        "Revenue": [100, 120, 90, 110, 130, 150, 95, 115, 105, 125, 140, 160],
        "Week": [1, 2, 1, 2, 1, 2, 1, 2, 1, 2, 1, 2],
        "Row_ID": list(range(12)),
    })
    assert _infer_factor_column(frame) == "Product"
    gate = _design_gate(frame, None, None, None)
    assert gate is None


def test_backend_auto_infers_generic_non_id_group_and_excludes_plot_patient_as_outcomes() -> None:
    from api import _design_gate, _infer_factor_column, _numeric_metric_columns

    frame = pd.DataFrame({
        "Plot": ["North", "North", "South", "South", "East", "East", "North", "North", "South", "South", "East", "East"],
        "Revenue": [100, 110, 120, 130, 140, 150, 105, 115, 125, 135, 145, 155],
        "Patient": ["P1", "P2", "P3", "P4", "P5", "P6", "P7", "P8", "P9", "P10", "P11", "P12"],
        "Shift": ["AM", "AM", "PM", "PM", "AM", "AM", "PM", "PM", "AM", "AM", "PM", "PM"],
        "Week": [1, 2, 1, 2, 1, 2, 1, 2, 1, 2, 1, 2],
    })
    assert _infer_factor_column(frame) == "Plot"
    assert _design_gate(frame, None, None, None) is None
    assert _numeric_metric_columns(frame, "Plot") == ["Revenue"]


def test_report_keeps_snake_case_labels_and_removes_forbidden_anova_phrase() -> None:
    from analysis_service import _discussion_limits, _humanize_label

    assert _humanize_label("hb_g_dl") == "hb_g_dl"
    text = _discussion_limits([
        {"test": "one-way anova", "groups": [{"name": "A", "n": 4}, {"name": "B", "n": 4}]}
    ])
    assert "ANOVA does not establish that every pair of groups differs" not in text
    assert "Post-hoc testing" in text


def test_docx_generation() -> None:
    frame = pd.DataFrame({"Treatment": ["A", "A", "B", "B"], "Value": [1, 2, 4, 5]})
    engine = analyze_dataframe(frame, "Treatment", "Value")
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "results.docx"
        assert generate_docx(engine, path) == str(path)
        assert path.is_file()


def test_engine_chart_is_returned_and_embedded_in_docx() -> None:
    frame = pd.DataFrame({"Treatment": ["A", "A", "B", "B"], "Value": [1, 2, 4, 5]})
    engine = analyze_dataframe(frame, "Treatment", "Value")
    chart_bytes = base64.b64decode(engine["chart_base64"])
    assert chart_bytes.startswith(b"\x89PNG\r\n\x1a\n")
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "results.docx"
        generate_docx(engine, path)
        with zipfile.ZipFile(path) as archive:
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
            assert media
            assert chart_bytes in [archive.read(name) for name in media]


def test_api_batch_returns_success_and_descriptive_fallbacks() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    text = """Record_ID,Teaching_Method,Exam_Score,ID_Number
R01,A,85.2,1
R02,A,88.5,1
R03,A,86.1,1
R04,B,72.4,2
R05,B,75.1,2
R06,B,71.8,2
R07,C,92.5,3
R08,C,94.1,3
R09,C,91.8,3
"""
    response = TestClient(app).post(
        "/api/v1/analyze",
        data={"raw_text": text, "metric_column": "ALL", "response_format": "json"},
        headers={"X-Rowfirst-Id": "demo-id", "X-Rowfirst-Secret-Key": "demo-secret"},
    )
    assert response.status_code == 200
    analyses = response.json()["analyses"]
    by_metric = {analysis["metric"]: analysis for analysis in analyses}
    assert by_metric["Exam_Score"]["status"] == "success"
    assert by_metric["Exam_Score"]["chart_base64"].startswith("iVBORw0KGgo")
    assert "ID_Number" not in by_metric


def test_api_gate_refuses_crm_row_id_and_index_design() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    text = "Company,Website,Revenue,Index,Row ID\n" + "\n".join(
        f"Company {index},example{index}.com,{100 + index},{index},{index}"
        for index in range(1, 101)
    )
    response = TestClient(app).post(
        "/api/v1/analyze",
        data={"raw_text": text, "metric_column": "ALL", "response_format": "json"},
        headers={"X-Rowfirst-Id": "demo-id", "X-Rowfirst-Secret-Key": "demo-secret"},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["mode"] in {"ask", "refuse"}
    assert payload["profile"]["n_rows"] == 100
    assert payload["reason"] in {"no_obvious_group_column", "unique_or_singleton_groups"}


def test_wide_company_like_table_requests_financial_schema_mapping() -> None:
    frame = pd.read_csv("stress_test_200col_1000rows.csv")
    from api import _design_gate

    response = _design_gate(frame, None, None, None)
    assert response is not None
    assert response["mode"] == "ask"
    assert response["reason"] == "incomplete_financial_schema"
    assert "date_or_month" in response["financial_schema"]["missing_required_fields"]
    assert "operating_expenses_or_costs" in response["financial_schema"]["missing_required_fields"]


def test_api_profile_mode_returns_value_counts_without_testing() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    response = TestClient(app).post(
        "/api/v1/analyze",
        data={
            "raw_text": "Department,Revenue\nEngineering,10\nEngineering,12\nSales,8\nSales,9\n",
            "factor_column": "Department",
            "mode": "profile",
            "response_format": "json",
        },
        headers={"X-Rowfirst-Id": "demo-id", "X-Rowfirst-Secret-Key": "demo-secret"},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["mode"] == "profile"
    assert payload["value_counts"] == {"Engineering": 2, "Sales": 2}


def test_api_batch_limits_charts_for_large_batches() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    headers = {"X-Rowfirst-Id": "demo-id", "X-Rowfirst-Secret-Key": "demo-secret"}
    rows = ["Treatment," + ",".join(f"Metric_{index}" for index in range(21))]
    for row_index, treatment in enumerate(("A", "B", "C") * 4):
        values = [str(row_index + index * 0.01) for index in range(21)]
        rows.append(f"{treatment}," + ",".join(values))
    response = TestClient(app).post(
        "/api/v1/analyze",
        data={"raw_text": "\n".join(rows), "metric_column": "ALL", "response_format": "json"},
        headers=headers,
    )
    assert response.status_code == 200
    analyses = response.json()["analyses"]
    assert len(analyses) == 21
    assert sum(analysis["chart_base64"] is not None for analysis in analyses) == 5


def test_large_batch_narrative_summarizes_non_significant_metrics() -> None:
    from analysis_service import _discussion, _interpretations

    significant = {
        "test": "one-way anova",
        "outcome": "Discovery",
        "p": 0.001,
        "isSignificant": True,
        "F": 12.0,
        "dfb": 2,
        "dfw": 9,
        "groups": [],
    }
    non_significant = [
        {"test": "one-way anova", "outcome": f"Metric_{index}", "p": 0.4, "isSignificant": False, "groups": []}
        for index in range(11)
    ]
    engine = {
        "ok": True,
        "results": [significant, *non_significant],
        "result": significant,
        "ingested": {"format": "long", "factor": "Treatment", "outcome": "Discovery", "groups": []},
    }
    interpretation = _interpretations(engine)[0]
    discussion = " ".join(_discussion(engine))
    assert "Discovery: group means were not reported" in interpretation
    assert "No significant differences were observed for 11 other variables tested" in interpretation
    assert "Metric_10: group means were not reported" not in interpretation
    assert "Metric_10: group means were not reported" not in discussion


def test_excel_and_pdf_uploads_are_parsed() -> None:
    excel_buffer = BytesIO()
    pd.DataFrame({"Treatment": ["A", "A", "B", "B"], "Value": [1, 2, 4, 5]}).to_excel(excel_buffer, index=False)
    excel = parse_excel_buffer(excel_buffer.getvalue(), "study.xlsx")
    assert list(excel.columns) == ["Treatment", "Value"]
    assert parse_uploaded_file(excel_buffer.getvalue(), "study.xlsx").shape == (4, 2)

    from reportlab.pdfgen.canvas import Canvas

    pdf_buffer = BytesIO()
    pdf = Canvas(pdf_buffer)
    for index, line in enumerate(("Treatment,Value", "A,1", "A,2", "B,4", "B,5")):
        pdf.drawString(72, 760 - index * 16, line)
    pdf.save()
    parsed_pdf = parse_pdf(pdf_buffer.getvalue())
    assert list(parsed_pdf.columns) == ["Treatment", "Value"]

    from docx import Document

    docx_buffer = BytesIO()
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Treatment"
    table.rows[0].cells[1].text = "Value"
    for treatment, value in (("A", "1"), ("B", "2")):
        cells = table.add_row().cells
        cells[0].text = treatment
        cells[1].text = value
    document.save(docx_buffer)
    docx = parse_docx_buffer(docx_buffer.getvalue())
    assert list(docx.columns) == ["Treatment", "Value"]
    assert parse_uploaded_file(docx_buffer.getvalue(), "study.docx").shape == (2, 2)


def test_image_upload_uses_ocr(monkeypatch) -> None:
    from PIL import Image
    import pytesseract

    image_buffer = BytesIO()
    Image.new("RGB", (120, 40), "white").save(image_buffer, format="PNG")
    monkeypatch.setattr(pytesseract, "image_to_string", lambda *args, **kwargs: "Treatment,Value\nA,1\nB,2")
    assert list(parse_image(image_buffer.getvalue()).columns) == ["Treatment", "Value"]


def test_bad_table_is_rejected() -> None:
    try:
        parse_tabular_text("single-value")
    except DataParserError:
        pass
    else:
        raise AssertionError("invalid delimiter was accepted")


def test_public_api_uses_shared_engine_and_docx_output() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    client = TestClient(app)
    headers = {"X-Rowfirst-ID": "demo-id", "X-Rowfirst-Secret": "demo-secret"}

    json_response = client.post(
        "/api/v1/analyze",
        data={
            "raw_text": "Treatment\tValue\nA\t1\nA\t2\nB\t4\nB\t5\n",
            "factor_column": "Treatment",
            "metric_column": "Value",
            "response_format": "json",
        },
        headers=headers,
    )
    assert json_response.status_code == 200
    payload = json_response.json()
    assert payload["ok"] is True
    assert payload["factor"] == "Treatment"
    assert payload["metric"] == "Value"
    accept_json_response = client.post(
        "/api/v1/analyze",
        data={
            "raw_text": "Treatment\tValue\nA\t1\nA\t2\nB\t4\nB\t5\n",
            "factor_column": "Treatment",
            "metric_column": "Value",
        },
        headers={**headers, "Accept": "application/json"},
    )
    assert accept_json_response.status_code == 200
    assert accept_json_response.headers["content-type"].startswith("application/json")
    assert accept_json_response.json()["breakdown"]

    docx_response = client.post(
        "/api/v1/analyze",
        data={
            "raw_text": "Treatment\tValue\nA\t1\nA\t2\nB\t4\nB\t5\n",
            "factor_column": "Treatment",
            "metric_column": "Value",
            "response_format": "docx",
        },
        headers=headers,
    )
    assert docx_response.status_code == 200
    assert docx_response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert len(docx_response.content) > 1000


def test_telegram_sends_verified_text_and_charts() -> None:
    from bot import _send_engine_outputs

    class Chat:
        id = 42

    class Message:
        chat = Chat()

    class RecordingBot:
        def __init__(self) -> None:
            self.messages = []
            self.photos = []

        def send_message(self, chat_id, text) -> None:
            self.messages.append((chat_id, text))

        def send_photo(self, chat_id, photo, caption) -> None:
            self.photos.append((chat_id, photo.read(), caption))

    bot = RecordingBot()
    chart = base64.b64encode(b"chart").decode("ascii")
    _send_engine_outputs(
        bot,
        Message(),
        {"breakdown": "Plain-English breakdown", "message": "Verified result", "chart_base64": chart},
    )

    assert [text for _, text in bot.messages] == [
        "Plain-English breakdown",
        "Verified results:\nVerified result",
    ]
    assert bot.photos == [(42, b"chart", "Monthly revenue and operating expenses")]


def test_public_api_accepts_form_query_and_bearer_credentials() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    client = TestClient(app)
    request_data = {
        "raw_text": "Treatment\tValue\nA\t1\nA\t2\nB\t4\nB\t5\n",
        "factor_column": "Treatment",
        "metric_column": "Value",
        "response_format": "json",
    }

    form_response = client.post(
        "/api/v1/analyze",
        data={**request_data, "rowfirst_id": "demo-id", "rowfirst_secret_key": "demo-secret"},
    )
    assert form_response.status_code == 200

    os.environ["ROWFIRST_ALLOW_QUERY_AUTH"] = "true"
    query_response = client.post(
        "/api/v1/analyze?rowfirst_id=demo-id&rowfirst_secret_key=demo-secret",
        data=request_data,
    )
    os.environ.pop("ROWFIRST_ALLOW_QUERY_AUTH", None)
    assert query_response.status_code == 200

    bearer_response = client.post(
        "/api/v1/analyze",
        data={**request_data, "rowfirst_id": "demo-id"},
        headers={"Authorization": "Bearer demo-secret"},
    )
    assert bearer_response.status_code == 200

    alias_header_response = client.post(
        "/api/v1/analyze",
        data=request_data,
        headers={"X-Rowfirst-Client-Id": "demo-id", "Rowfirst-Secret-Key": "demo-secret"},
    )
    assert alias_header_response.status_code == 200


def test_public_api_infers_columns_when_not_supplied() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    response = TestClient(app).post(
        "/api/v1/analyze",
        data={
            "raw_text": "Treatment,Value\nA,1\nA,2\nA,3\nB,4\nB,5\nB,6\n",
            "response_format": "json",
        },
        headers={"X-Rowfirst-Id": "demo-id", "X-Rowfirst-Secret-Key": "demo-secret"},
    )
    assert response.status_code == 200
    assert response.json()["factor"] == "Treatment"
    assert response.json()["metric"] == "Value"


def test_public_api_supports_explicit_advanced_models_and_docx() -> None:
    os.environ["ROWFIRST_ID"] = "demo-id"
    os.environ["ROWFIRST_SECRET_KEY"] = "demo-secret"
    frame = pd.DataFrame({
        "X1": [0, 0, 1, 1, 2, 0, 2, 1, 3, 0],
        "X2": [0, 1, 0, 1, 0, 2, 1, 2, 0, 3],
        "Y": [1, 4, 3, 6, 5, 7, 8, 9, 7, 10],
        "Outcome": [0, 0, 0, 1, 1, 1, 1, 1, 1, 1],
        "Month": list(range(1, 11)),
    })
    raw_text = frame.to_csv(index=False)
    headers = {
        "X-Rowfirst-Id": "demo-id",
        "X-Rowfirst-Secret-Key": "demo-secret",
        "Accept": "application/json",
    }
    client = TestClient(app)

    multiple = client.post(
        "/api/v1/analyze",
        data={"raw_text": raw_text, "mode": "multiple_regression", "predictor_columns": "X1,X2", "outcome_column": "Y"},
        headers=headers,
    )
    assert multiple.status_code == 200
    assert multiple.json()["result"]["test"] == "multiple linear regression"

    logistic = client.post(
        "/api/v1/analyze",
        data={"raw_text": raw_text, "mode": "logistic_regression", "predictor_columns": "X1,X2", "outcome_column": "Outcome"},
        headers=headers,
    )
    assert logistic.status_code == 200
    assert logistic.json()["result"]["test"] == "logistic regression"

    forecast = client.post(
        "/api/v1/analyze",
        data={"raw_text": raw_text, "mode": "forecast", "outcome_column": "Y", "time_column": "Month", "horizon": "2"},
        headers=headers,
    )
    assert forecast.status_code == 200
    assert forecast.json()["result"]["test"] == "linear forecast"
    assert len(forecast.json()["result"]["forecast"]) == 2

    docx_response = client.post(
        "/api/v1/analyze",
        data={"raw_text": raw_text, "mode": "multiple_regression", "predictor_columns": "X1,X2", "outcome_column": "Y", "response_format": "docx"},
        headers={key: value for key, value in headers.items() if key != "Accept"},
    )
    assert docx_response.status_code == 200
    assert docx_response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
