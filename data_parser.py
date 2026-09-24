"""Framework-free tabular input normalization."""
from __future__ import annotations

import base64
import csv
import json
import mimetypes
import os
import re
from io import BytesIO, StringIO
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urlparse, urlunparse
from urllib.request import Request, urlopen

import pandas as pd


class DataParserError(ValueError):
    """Raised when tabular input cannot be normalized."""


def parse_csv_buffer(buffer: bytes | bytearray | BytesIO, filename: str = "") -> pd.DataFrame:
    """Parse a CSV upload buffer into a DataFrame."""
    if isinstance(buffer, BytesIO):
        payload = buffer.getvalue()
    else:
        payload = bytes(buffer)
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DataParserError("CSV files must be UTF-8 encoded") from exc
    if filename and Path(filename).suffix.lower() != ".csv":
        raise DataParserError("Only .csv uploads are supported")
    return parse_tabular_text(text)


def parse_excel_buffer(buffer: bytes | bytearray | BytesIO, filename: str = "") -> pd.DataFrame:
    """Read the first worksheet from an XLSX or XLS upload."""
    payload = buffer.getvalue() if isinstance(buffer, BytesIO) else bytes(buffer)
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xls"}:
        raise DataParserError("Only .xlsx and .xls uploads are supported")
    try:
        frame = pd.read_excel(BytesIO(payload), sheet_name=0)
    except (ImportError, ValueError, OSError) as exc:
        raise DataParserError(f"Could not read Excel file: {exc}") from exc
    return _clean_frame(frame, "Excel worksheet")


def parse_pdf(buffer: bytes | bytearray | BytesIO, filename: str = "", content_type: str | None = None) -> pd.DataFrame:
    """Extract tabular text from a PDF, preferring Mistral OCR if configured."""
    payload = buffer.getvalue() if isinstance(buffer, BytesIO) else bytes(buffer)
    try:
        if _mistral_ocr_enabled():
            return extract_mistral_table_frame(_mistral_ocr_extract(payload, filename or "upload.pdf", content_type=content_type))
    except DataParserError:
        pass
    try:
        from pypdf import PdfReader

        text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(payload)).pages)
    except Exception as exc:
        raise DataParserError(f"Could not read PDF file: {exc}") from exc
    if text.strip():
        return parse_tabular_text(text)
    try:
        import fitz
        from PIL import Image
        import pytesseract

        pages = []
        document = fitz.open(stream=payload, filetype="pdf")
        for page in document:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
            pages.append(pytesseract.image_to_string(image, config="--psm 6"))
        text = "\n".join(pages)
    except Exception as exc:
        raise DataParserError(f"PDF contains no text and OCR is unavailable: {exc}") from exc
    return parse_tabular_text(text)


def parse_docx_buffer(buffer: bytes | bytearray | BytesIO) -> pd.DataFrame:
    """Read the first Word table, or tabular text from Word paragraphs."""
    payload = buffer.getvalue() if isinstance(buffer, BytesIO) else bytes(buffer)
    try:
        from docx import Document

        document = Document(BytesIO(payload))
    except Exception as exc:
        raise DataParserError(f"Could not read Word file: {exc}") from exc
    if document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in document.tables[0].rows]
        if len(rows) < 2:
            raise DataParserError("Word table must include headers and data rows")
        return _clean_frame(pd.DataFrame(rows[1:], columns=rows[0]), "Word table")
    text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    if not text:
        raise DataParserError("Word file contains no tabular data")
    return parse_tabular_text(text)


def parse_image(buffer: bytes | bytearray | BytesIO, filename: str = "", content_type: str | None = None) -> pd.DataFrame:
    """OCR a table image and normalize the recognized text, preferring Mistral OCR if configured."""
    payload = buffer.getvalue() if isinstance(buffer, BytesIO) else bytes(buffer)
    try:
        if _mistral_ocr_enabled():
            return extract_mistral_table_frame(_mistral_ocr_extract(payload, filename or "upload.png", content_type=content_type))
    except DataParserError:
        pass
    try:
        from PIL import Image, ImageOps
        import pytesseract

        image = Image.open(BytesIO(payload)).convert("L")
        image = ImageOps.autocontrast(image)
        text = pytesseract.image_to_string(image, config="--psm 6")
    except Exception as exc:
        raise DataParserError(f"Could not OCR image: {exc}") from exc
    return parse_tabular_text(text)


def detect_upload_kind(filename: str = "", content_type: str | None = None, payload: bytes | bytearray | BytesIO | None = None) -> str:
    """Infer the upload kind from MIME type and file extension while tolerating missing suffixes."""
    normalized_name = (filename or "").lower()
    normalized_type = (content_type or "").lower()
    suffix = Path(normalized_name).suffix.lower()

    if "csv" in normalized_type or suffix == ".csv":
        return "csv"
    if "excel" in normalized_type or suffix in {".xlsx", ".xls"}:
        return "excel"
    if "pdf" in normalized_type or suffix == ".pdf":
        return "pdf"
    if "word" in normalized_type or suffix == ".docx":
        return "docx"
    if "image/" in normalized_type or suffix in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}:
        return "image"

    if payload is not None:
        sample = payload.getvalue() if isinstance(payload, BytesIO) else bytes(payload)
        if sample.startswith(b"%PDF"):
            return "pdf"
        if sample.startswith((b"\x89PNG", b"\xff\xd8\xff", b"GIF8", b"II*\x00", b"BM")):
            return "image"

    raise DataParserError("Unsupported upload type; expected CSV, image, PDF, DOCX, or spreadsheet")


def parse_uploaded_file(buffer: bytes | bytearray | BytesIO, filename: str, content_type: str | None = None) -> pd.DataFrame:
    """Dispatch supported uploads to CSV, Excel, PDF, Word, or image parsing."""
    payload = buffer.getvalue() if isinstance(buffer, BytesIO) else bytes(buffer)
    kind = detect_upload_kind(filename, content_type, payload)
    if kind == "csv":
        return parse_csv_buffer(buffer, filename)
    if kind == "excel":
        return parse_excel_buffer(buffer, filename)
    if kind == "pdf":
        return parse_pdf(buffer, filename=filename, content_type=content_type)
    if kind == "docx":
        return parse_docx_buffer(buffer)
    if kind == "image":
        return parse_image(buffer, filename=filename, content_type=content_type)
    raise DataParserError("Supported uploads are CSV, XLSX, XLS, PDF, DOCX, PNG, JPG, WEBP, TIFF, or BMP")


def parse_spreadsheet_source(source: str) -> pd.DataFrame:
    """Load a spreadsheet from a Google Sheets URL or a remote spreadsheet file URL."""
    text = str(source or "").strip()
    if not text:
        raise DataParserError("Spreadsheet source is empty")
    parsed = urlparse(text)
    if parsed.scheme in {"http", "https"}:
        path = parsed.path.lower()
        suffix = Path(parsed.path).suffix.lower()
        is_google_sheet = (
            "docs.google.com" in parsed.netloc.lower() and "/spreadsheets/" in path
        ) or "spreadsheets.google.com" in parsed.netloc.lower()
        if is_google_sheet:
            export_url = _google_sheet_export_url(text)
            payload = _download_url_bytes(export_url)
            return _parse_remote_spreadsheet_payload(payload, "google_sheet.csv", source_hint="google_sheets")
        if suffix in {".csv", ".xlsx", ".xls"}:
            payload = _download_url_bytes(text)
            filename = Path(parsed.path).name or "remote_spreadsheet"
            return parse_uploaded_file(payload, filename)
        if "format=csv" in text.lower() or "export?format=" in text.lower() or text.lower().endswith("csv"):
            payload = _download_url_bytes(text)
            return _parse_remote_spreadsheet_payload(payload, Path(parsed.path).name or "remote_data.csv", source_hint="csv_export")
        if "format=xlsx" in text.lower() or "format=xls" in text.lower():
            payload = _download_url_bytes(text)
            filename = Path(parsed.path).name or "remote_spreadsheet.xlsx"
            return parse_uploaded_file(payload, filename)
    raise DataParserError("Only Google Sheets URLs and remote CSV/XLSX/XLS spreadsheet URLs are supported")


def _parse_remote_spreadsheet_payload(payload: bytes, filename: str, source_hint: str = "remote") -> pd.DataFrame:
    """Parse remote spreadsheet payloads safely, falling back from CSV to Excel when needed."""
    try:
        return parse_csv_buffer(payload, filename=filename)
    except DataParserError as exc:
        if payload.startswith(b"PK"):
            try:
                return parse_excel_buffer(payload, filename=filename if filename.lower().endswith((".xlsx", ".xls")) else "remote_spreadsheet.xlsx")
            except DataParserError:
                raise exc from exc
        raise exc from exc


def _google_sheet_export_url(sheet_url: str) -> str:
    url = urlparse(sheet_url)
    host = url.netloc.lower()
    path = url.path.lower()
    if "docs.google.com" not in host and "spreadsheets.google.com" not in host:
        raise DataParserError("Not a valid Google Sheets URL")
    if "/d/" not in path and "/spreadsheets/d/" not in path:
        raise DataParserError("Google Sheets URL must include a spreadsheet id")
    spreadsheet_id = path.split("/d/", 1)[1].split("/", 1)[0]
    query = parse_qs(url.query)
    requested_format = (query.get("format") or [None])[0]
    if requested_format in {"xlsx", "xls"}:
        return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?{urlencode({'format': requested_format})}"
    return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export"


def _download_url_bytes(url: str) -> bytes:
    try:
        with urlopen(url, timeout=30) as response:
            return response.read()
    except Exception as exc:  # pragma: no cover - runtime guard
        raise DataParserError(f"Could not download spreadsheet source: {exc}") from exc


def parse_tabular_text(text: str) -> pd.DataFrame:
    """Parse comma-, tab-, pipe-delimited markdown, or whitespace-delimited tabular text."""
    cleaned = str(text or "").strip()
    if not cleaned:
        raise DataParserError("No tabular data was supplied")

    lines = [line for line in cleaned.splitlines() if line.strip()]
    labeled = _parse_labeled_rows(lines)
    if labeled is not None:
        return labeled

    if any("|" in line for line in lines):
        return _parse_markdown_table(lines)

    delimiter = _detect_delimiter(lines)
    try:
        if delimiter == "whitespace":
            frame = pd.read_csv(StringIO("\n".join(lines)), sep=r"\s+", engine="python")
        else:
            frame = pd.read_csv(StringIO("\n".join(lines)), sep=delimiter)
    except (csv.Error, pd.errors.ParserError, ValueError) as exc:
        raise DataParserError("Could not determine a valid table delimiter") from exc
    frame.columns = [str(column).strip() for column in frame.columns]
    if frame.empty or any(not column for column in frame.columns):
        raise DataParserError("The tabular input must include headers and data rows")
    return frame


def _parse_markdown_table(lines: list[str]) -> pd.DataFrame:
    clean_lines = []
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("```"):
            continue
        clean_lines.append(stripped.strip("|"))
    if not clean_lines:
        raise DataParserError("No markdown table data was supplied")

    row_values = [
        [cell.strip() for cell in line.split("|") if cell.strip()]
        for line in clean_lines
    ]
    if len(row_values) < 2:
        raise DataParserError("Markdown tables must include headers and data rows")

    header = row_values[0]
    body = row_values[1:]
    if all(re.fullmatch(r":?-{3,}:?", cell) for cell in row_values[1]):
        body = row_values[2:]
    if not body:
        raise DataParserError("Markdown table must contain data rows")
    frame = pd.DataFrame(body, columns=header)
    if frame.empty or any(not column for column in frame.columns):
        raise DataParserError("The markdown table must include non-empty headers")
    return frame


def _detect_delimiter(lines: list[str]) -> str:
    sample = "\n".join(lines[:10])
    if "\t" in sample:
        return "\t"
    if "|" in sample:
        return "|"
    if "," in sample:
        return ","
    if ";" in sample:
        return ";"
    if any(len(line.split()) > 1 for line in lines):
        return "whitespace"
    raise DataParserError("Could not determine a valid table delimiter")


def _parse_labeled_rows(lines: list[str]) -> pd.DataFrame | None:
    if not all(":" in line for line in lines):
        return None
    rows: list[dict[str, object]] = []
    for line in lines:
        label, values = line.split(":", 1)
        parsed_values = [value.strip() for value in values.split(",") if value.strip()]
        if not label.strip() or not parsed_values:
            raise DataParserError("Labeled rows must contain a name and values")
        for value in parsed_values:
            rows.append({"Factor": label.strip(), "Metric": value})
    return pd.DataFrame(rows)


def _clean_frame(frame: pd.DataFrame, source: str) -> pd.DataFrame:
    frame = frame.dropna(axis=0, how="all").dropna(axis=1, how="all")
    if frame.empty:
        raise DataParserError(f"{source} is empty")
    frame.columns = [str(column).strip() for column in frame.columns]
    if any(not column for column in frame.columns):
        raise DataParserError(f"{source} must include non-empty headers")
    return frame


def _mistral_ocr_enabled() -> bool:
    return bool(os.getenv("MISTRAL_API_KEY"))


def _mistral_ocr_extract(payload: bytes, filename: str, content_type: str | None = None) -> dict | list | str:
    api_key = os.getenv("MISTRAL_API_KEY")
    if not api_key:
        raise DataParserError("Mistral API key is not configured")

    endpoint = os.getenv("MISTRAL_OCR_ENDPOINT", "https://api.mistral.ai/v1/ocr")
    model = os.getenv("MISTRAL_MODEL", "mistral-ocr-latest")
    body = {
        "model": model,
        "document": {
            "type": "file",
            "file_name": filename or "upload",
            "content": base64.b64encode(payload).decode("utf-8"),
        },
    }
    if content_type:
        body["document"]["mime_type"] = content_type

    request = Request(
        endpoint,
        data=json.dumps(body).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "User-Agent": "rowfirst/1.0",
        },
        method="POST",
    )
    try:
        with urlopen(request, timeout=60) as response:
            response_text = response.read().decode("utf-8")
    except Exception as exc:
        raise DataParserError(f"Could not access Mistral OCR API: {exc}") from exc

    try:
        return json.loads(response_text)
    except json.JSONDecodeError:
        return response_text


def extract_mistral_table_frame(extracted: object) -> pd.DataFrame:
    """Normalize Mistral OCR output into a pandas DataFrame."""
    if isinstance(extracted, str):
        try:
            parsed = json.loads(extracted)
        except json.JSONDecodeError:
            parsed = extracted
    else:
        parsed = extracted

    if isinstance(parsed, dict):
        for key in ("markdown", "md", "table", "table_markdown", "content", "result"):
            if key in parsed and parsed[key] not in (None, ""):
                return extract_mistral_table_frame(parsed[key])
        for key in ("rows", "data"):
            if key in parsed and parsed[key] is not None:
                return pd.DataFrame(parsed[key])
        if "columns" in parsed and "data" in parsed:
            return pd.DataFrame(parsed["data"], columns=parsed["columns"])
        if "pages" in parsed:
            return extract_mistral_table_frame(parsed["pages"])
        if "tables" in parsed:
            return extract_mistral_table_frame(parsed["tables"])
        if parsed and all(isinstance(value, list) for value in parsed.values()):
            return pd.DataFrame(parsed)

    if isinstance(parsed, list):
        if not parsed:
            raise DataParserError("Mistral returned no table data")
        if all(isinstance(item, dict) for item in parsed):
            nested_values = []
            for item in parsed:
                nested_values.append({_k: _v for _k, _v in item.items() if _k not in {"markdown", "table"}})
            if nested_values and all(set(item) == set(nested_values[0]) for item in nested_values):
                return pd.DataFrame(nested_values)
        if all(isinstance(item, (list, tuple)) for item in parsed):
            return pd.DataFrame(parsed)
        for item in parsed:
            try:
                return extract_mistral_table_frame(item)
            except DataParserError:
                continue

    if isinstance(parsed, str):
        cleaned = parsed.strip()
        if cleaned.startswith("```"):
            cleaned = re.sub(r"^```(?:markdown|json)?\s*|```\s*$", "", cleaned, flags=re.IGNORECASE)
        if cleaned.startswith("[") or cleaned.startswith("{"):
            try:
                return extract_mistral_table_frame(json.loads(cleaned))
            except json.JSONDecodeError:
                pass
        return parse_tabular_text(cleaned)

    raise DataParserError("Mistral OCR response did not include a valid table")